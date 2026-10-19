"""Excel workbooks, values from the text after the table, and the column operations
(TGA % of initial mass, DSC heat flow per sample mass)."""

import openpyxl
import pandas as pd
import pytest

from mfethuls.instruments import workbook
from mfethuls.instruments.definition import DefinitionError, parse_definition
from mfethuls.instruments.detect import detect, file_outline
from mfethuls.instruments.spec import json_schema, spec_card
from mfethuls.parsers.tabular import TabularParser, apply_derived


def _parse(definition, *paths):
    return TabularParser(definition).parse({"x": [str(p) for p in paths]}, experiment_id="e1",
                                           instrument_type=definition.type)


def _tensile_workbook(path, sample="LUB001_test", rows=((0.0, 0.001), (0.01, 0.002), (0.02, 0.004))):
    book = openpyxl.Workbook()
    params = book.active
    params.title = "Parameters"
    params.append(["Test standard", None, "EN 2243-1"])
    results = book.create_sheet("Results")
    results.append([None, "Shearing strength", "Maximum force"])
    results.append([None, "MPa", "N"])
    results.append([sample, 0.0052, 0.84])
    data = book.create_sheet(sample[:20])  # named after the sample, as the instrument does
    data.append([sample, sample])
    data.append(["Strain", "Shearing stress"])
    data.append(["mm", "MPa"])
    for row in rows:
        data.append(list(row))
    book.save(path)
    return path


TENSILE = {
    "name": "tensile_tester", "type": "tensile", "model": "zwick", "folder_name": "tensile_tester",
    "file": {"extensions": [".xlsx"], "header": {"starts_at": "Strain", "rows": 2}},
    "columns": [
        {"source": "Strain mm", "quantity": "strain", "unit": "mm"},
        {"source": "Shearing stress MPa", "quantity": "shear_stress", "unit": "MPa"},
    ],
}


# --- workbooks ----------------------------------------------------------------

def test_workbook_table_is_read_from_the_sheet_with_the_header(tmp_path):
    path = _tensile_workbook(tmp_path / "LUB001_test.xlsx")
    dataset = _parse(parse_definition(TENSILE), path)
    data = dataset.data
    assert data["strain_mm"].tolist() == [0.0, 0.01, 0.02]
    assert data["shear_stress_mpa"].tolist() == [0.001, 0.002, 0.004]  # MPa, the convention for tensile tests
    text = dataset.metadata["file_text"]["LUB001_test.xlsx"]
    assert text["sheets"] == ["LUB001_test"]
    assert data["sheet"].unique().tolist() == ["LUB001_test"]
    # Nothing is dropped: the other sheets are kept as text.
    assert "sheet: Results" in text and "sheet: Parameters" in text
    assert any("0.0052" in line for line in text["sheet: Results"])


def test_workbook_sheet_by_number_and_values_from_other_sheets(tmp_path):
    path = _tensile_workbook(tmp_path / "LUB001_test.xlsx")
    definition = parse_definition({
        **TENSILE,
        "file": {**TENSILE["file"], "sheet": -1},
        "from_header": [{"column": "max_force_N", "pattern": r"^LUB\S*\t[^\t]*\t([0-9.]+)", "type": "number"}],
    })
    dataset = _parse(definition, path)
    assert dataset.data["max_force_N"].tolist() == [0.84] * 3
    assert dataset.metadata["file_text"]["LUB001_test.xlsx"]["values"] == {"max_force_N": 0.84}


@pytest.mark.parametrize("file_settings,problem", [
    ({"extensions": [".xlsx"], "delimiter": "comma"}, "tab-separated"),
    ({"extensions": [".xlsx"], "decimal": ","}, "hold numbers"),
    ({"extensions": [".txt"], "sheet": 2}, "only goes with workbooks"),
    ({"extensions": [".xlsx"], "sheet": 0}, "1 = first"),
    ({"extensions": [".xlsx", ".csv"]}, "mixes workbooks"),
])
def test_workbook_settings_are_checked(file_settings, problem):
    with pytest.raises(DefinitionError) as info:
        parse_definition({**TENSILE, "file": file_settings})
    assert problem in str(info.value)


def test_pick_sheet():
    sheets = {"Parameters": ["a"], "Results": ["MPa"], "LUB": ["Strain", "mm\tMPa"]}
    assert workbook.pick_sheet(sheets, None, "Strain") == "LUB"
    assert workbook.pick_sheet(sheets, None, None) == "Parameters"
    assert workbook.pick_sheet(sheets, -1) == "LUB"
    assert workbook.pick_sheet(sheets, 2) == "Results"
    with pytest.raises(ValueError, match="no sheet 5"):
        workbook.pick_sheet(sheets, 5)
    with pytest.raises(ValueError, match="no sheet 'Data'"):
        workbook.pick_sheet(sheets, "Data")


def test_detect_reads_a_workbook_with_bare_units(tmp_path):
    path = _tensile_workbook(tmp_path / "LUB001_test.xlsx")
    detection = detect(path)
    assert detection.file["extensions"] == [".xlsx"]
    assert "delimiter" not in detection.file and "encoding" not in detection.file
    assert detection.file["header"] == {"rows": 2, "starts_at": "Strain"}
    guesses = {g.source: (g.quantity, g.unit) for g in detection.columns}
    assert guesses == {"Strain mm": ("strain", "mm"), "Shearing stress MPa": ("shear_stress", "MPa")}
    definition = parse_definition(detection.definition(name="t", type="tensile", model="m", folder_name="f"))
    assert len(_parse(definition, path).data) == 3


def test_prompt_shows_workbook_sheets_not_bytes(tmp_path):
    path = _tensile_workbook(tmp_path / "LUB001_test.xlsx")
    card = spec_card(path)
    assert "Sheet 3: 'LUB001_test' (6 rows)" in card
    assert "Strain\tShearing stress" in card
    assert "PK" not in card.split("The Excel workbook")[1]
    assert "sheet" in json_schema()["properties"]["file"]["properties"]
    assert file_outline(path)


def test_binary_files_are_refused_with_a_message(tmp_path):
    path = tmp_path / "run.ds8d"
    path.write_bytes(b"\x00\x01binary\x00data" * 50)
    with pytest.raises(ValueError, match="not a text file"):
        spec_card(path)
    with pytest.raises(ValueError, match="not a text file"):
        detect(path)


# --- values after the table; column operations ----------------------------------

DSC_PRIOR_LIKE = """Curve Values:
 Index t Value
 [#] [s] [Wg^-1]
 0 0 -0.5
 1 1 -1.0
Sample:
  BKBMBTT, 5.0000 mg
"""


def test_sample_mass_from_the_footer_gives_heat_flow_in_mW(tmp_path):
    path = tmp_path / "lub1.txt"
    path.write_text(DSC_PRIOR_LIKE, encoding="utf-8")
    definition = parse_definition({
        "name": "dsc_x", "type": "dsc", "model": "x", "folder_name": "DSC",
        "file": {"extensions": [".txt"], "delimiter": "whitespace", "header": {"starts_at": "Index", "rows": 2},
                 "data_ends": "first_non_numeric_row"},
        "columns": [{"source": "t [s]", "quantity": "time", "unit": "s"},
                    {"source": "Value [Wg^-1]", "quantity": "specific_heat_flow", "unit": "W/g"}],
        "from_header": [{"column": "sample_mass_mg", "pattern": r"^\s*\S.*,\s*([0-9.]+)\s*mg\s*$",
                         "type": "number"}],
        "derived": [{"op": "multiply", "column": "heat_flow_W_g", "by": "sample_mass_mg", "into": "heat_flow_mW"}],
    })
    data = _parse(definition, path).data
    assert data["sample_mass_mg"].tolist() == [5.0, 5.0]
    assert data["heat_flow_W_g"].tolist() == [-0.5, -1.0]  # kept
    assert data["heat_flow_mW"].tolist() == [-2.5, -5.0]


def test_divide_skips_files_without_the_other_column():
    definition = parse_definition({
        "name": "d", "type": "dsc", "model": "x", "folder_name": "DSC", "file": {"extensions": [".txt"]},
        "columns": [{"source": "HF", "quantity": "heat_flow", "unit": "mW"}],
        "derived": [{"op": "divide", "column": "heat_flow_mW", "by": "sample_mass_mg", "into": "heat_flow_W_g"}],
    })
    df = pd.DataFrame({"heat_flow_mW": [1.0, 2.0]})
    out, warnings = apply_derived(df.copy(), definition)
    assert "heat_flow_W_g" not in out.columns and not warnings
    out, _ = apply_derived(df.assign(sample_mass_mg=[4.0, 0.0]), definition)
    assert out["heat_flow_W_g"].iloc[0] == 0.25 and pd.isna(out["heat_flow_W_g"].iloc[1])  # no division by 0


@pytest.mark.parametrize("step,problem", [
    ({"op": "divide", "column": "a", "into": "b"}, "needs 'by'"),
    ({"op": "scale", "column": "a", "into": "b", "factor": 2, "by": "c"}, "'by' only goes with"),
])
def test_column_operations_are_checked(step, problem):
    with pytest.raises(DefinitionError) as info:
        parse_definition({"name": "d", "type": "dsc", "model": "x", "folder_name": "DSC",
                          "file": {"extensions": [".txt"]},
                          "columns": [{"source": "HF", "quantity": "heat_flow", "unit": "mW"}], "derived": [step]})
    assert problem in str(info.value)


def test_pct_of_initial_starts_each_curve_at_100():
    definition = parse_definition({
        "name": "t", "type": "tga", "model": "x", "folder_name": "TGA", "file": {"extensions": [".txt"]},
        "columns": [{"source": "m", "quantity": "mass", "unit": "mg"}],
        "derived": [{"op": "pct_of_initial", "column": "mass_mg", "into": "mass_pct"}],
    })
    df = pd.DataFrame({"name": ["a", "a", "b", "b"], "mass_mg": [10.0, 5.0, 2.0, 2.5]})
    out, _ = apply_derived(df, definition)
    assert out["mass_pct"].tolist() == [100.0, 50.0, 100.0, 125.0]


# --- one quantity, several contexts -----------------------------------------------

def test_strain_in_mm_and_in_percent_never_share_a_column():
    from mfethuls.instruments.units import QUANTITIES, conversion

    strain = QUANTITIES["strain"]
    assert strain.target("rheometer", "%") == ("strain_pct", "%")
    assert strain.target("rheometer", "fraction") == ("strain_pct", "%")
    assert strain.target("tensile", "mm") == ("strain_mm", "mm")
    assert strain.target("dma") == ("strain_pct", "%")
    assert conversion(strain, "um") == (0.001, 0.0)  # into mm, never into %
    assert conversion(strain, "fraction") == (100.0, 0.0)


def test_shear_stress_is_mpa_for_tensile_and_pa_for_rheology():
    from mfethuls.instruments.units import QUANTITIES, conversion

    stress = QUANTITIES["shear_stress"]
    assert stress.target("tensile") == ("shear_stress_mpa", "MPa")
    assert stress.target("rheometer") == ("shear_stress_pa", "Pa")
    assert conversion(stress, "MPa", "tensile") == (1.0, 0.0)
    assert conversion(stress, "MPa", "rheometer") == (1e6, 0.0)


# --- several tests in one workbook ------------------------------------------------

def _multi_test_workbook(path, specimens=("LUB640", "LUB640_1", "LUB641")):
    book = openpyxl.Workbook()
    book.active.title = "Parameters"
    book.active.append(["Test standard", None, "ASTM D 638"])
    results = book.create_sheet("Results")
    results.append([None, "sM", "eB"])
    results.append([None, "kPa", "%"])
    for i, specimen in enumerate(specimens):
        results.append([specimen, 20.0 + i, 300.0 + i])
        sheet = book.create_sheet(specimen)
        sheet.append([specimen, specimen])
        sheet.append(["Elongation", "Standard force"])
        sheet.append(["%", "kPa"])
        for j in range(3):
            sheet.append([float(j), 100.0 * (j + i)])
    book.save(path)
    return path


def test_every_sheet_with_the_table_is_read_one_curve_per_sheet(tmp_path):
    path = _multi_test_workbook(tmp_path / "LUB640toLUB641.xlsx")
    definition = parse_definition({
        "name": "tensile_tester", "type": "tensile", "model": "zwick", "folder_name": "tensile_tester",
        "name_from": "sheet",
        "file": {"extensions": [".xlsx"], "header": {"starts_at": "Elongation", "rows": 2}},
        "columns": [{"source": "Elongation %", "quantity": "strain", "unit": "%"},
                    {"source": "Standard force kPa", "quantity": "stress", "unit": "kPa"}],
    })
    dataset = _parse(definition, path)
    data = dataset.data
    assert data.groupby("name").size().to_dict() == {"LUB640": 3, "LUB640_1": 3, "LUB641": 3}  # nothing skipped
    assert data["strain_pct"].tolist()[:3] == [0.0, 1.0, 2.0]
    assert data.loc[data["name"] == "LUB641", "stress_mpa"].tolist() == pytest.approx([0.2, 0.3, 0.4])  # kPa -> MPa
    text = dataset.metadata["file_text"]["LUB640toLUB641.xlsx"]
    assert text["sheets"] == ["LUB640", "LUB640_1", "LUB641"]
    assert "sheet: Results" in text and "preamble (LUB641)" in text


def test_detect_proposes_one_curve_per_sheet(tmp_path):
    path = _multi_test_workbook(tmp_path / "LUB640toLUB641.xlsx")
    detection = detect(path)
    assert detection.name_from == "sheet" and "sheet" not in detection.file
    assert {g.source: (g.quantity, g.unit) for g in detection.columns} == {
        "Elongation %": ("strain", "%"), "Standard force kPa": ("stress", "kPa")}
    draft = detection.definition(name="t", type="tensile", model="m", folder_name="f")
    assert draft["name_from"] == "sheet"
    assert set(_parse(parse_definition(draft), path).data["name"]) == {"LUB640", "LUB640_1", "LUB641"}


def test_a_single_test_sheet_keeps_the_file_name(tmp_path):
    """Sheets named after the sample are often cut short ('LUBxxx_80C_DCMGel...'); one test is named after its file."""
    path = _tensile_workbook(tmp_path / "LUB207_80C_DCMGel_10min.xlsx", sample="LUBxxx_80C_DCMGel_10min_long")
    data = _parse(parse_definition({**TENSILE, "name_from": "sheet"}), path).data
    assert set(data["name"]) == {"LUB207_80C_DCMGel_10min"}
    assert set(data["sheet"]) == {"LUBxxx_80C_DCMGel_10"}


def test_each_test_gets_its_date_from_its_results_row(tmp_path):
    path = tmp_path / "LUB640toLUB641.xlsx"
    book = openpyxl.Workbook()
    results = book.active
    results.title = "Results"
    results.append([None, "Date/Clock time", "L0"])
    results.append(["LUB640", 46246.745150462964, 8.8])  # Excel day number in a 'General' cell
    results.append(["LUB640_1", 46246.75, 9.1])
    for specimen in ("LUB640", "LUB640_1"):
        sheet = book.create_sheet(specimen)
        sheet.append([specimen, specimen])
        sheet.append(["Elongation", "Standard force"])
        sheet.append(["%", "kPa"])
        sheet.append([0.0, 1.0])
        sheet.append([1.0, 2.0])
    book.save(path)
    definition = parse_definition({
        "name": "t", "type": "tensile", "model": "m", "folder_name": "f", "name_from": "sheet",
        "file": {"extensions": [".xlsx"], "header": {"starts_at": "Elongation", "rows": 2}},
        "columns": [{"source": "Elongation %", "quantity": "strain", "unit": "%"}],
        "from_header": [{"column": "test_date", "pattern": "^{sheet}\t([0-9.]+)", "type": "datetime",
                         "timezone": "Europe/Zurich"}],
    })
    dataset = _parse(definition, path)
    dates = dataset.data.groupby("name")["test_date"].first()
    # 17:53:01 summer time in Zurich; "LUB640" must not pick up the row of "LUB640_1".
    assert str(dates["LUB640"]) == "2026-08-12 15:53:01+00:00"
    assert str(dates["LUB640_1"]) == "2026-08-12 16:00:00+00:00"
    assert dataset.metadata["file_text"][path.name]["values"]["test_date"] == {
        "LUB640": "2026-08-12T15:53:01+00:00", "LUB640_1": "2026-08-12T16:00:00+00:00"}
