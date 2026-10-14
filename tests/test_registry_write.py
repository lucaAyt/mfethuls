"""Appending new rows to the registry file without disturbing what is there."""

import os

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from mfethuls.experiments import append_rows

NEW_ROWS = pd.DataFrame(
    [
        {"name": "poly9", "instrument_name": "tga", "sample_id": "S009", "run_id": "R001",
         "raw_data_filename": "poly9", "batch": "B2", "_folder": "TGA", "_modified": None},
    ]
)


def test_csv_append_keeps_existing_text_adds_columns_and_backs_up(tmp_path):
    path = tmp_path / "registry.csv"
    original = b"name,instrument_name,sample_id,run_id,mass_mg\nLB_tga_001,tga,S001,R001,10\n"
    path.write_bytes(original)

    registry = append_rows(str(path), NEW_ROWS)

    assert path.read_bytes() == (
        b"name,instrument_name,sample_id,run_id,mass_mg,raw_data_filename,batch\n"
        b"LB_tga_001,tga,S001,R001,10,,\n"  # "10" stays "10", not "10.0"
        b"poly9,tga,S009,R001,,poly9,B2\n"
    )
    assert (tmp_path / "registry.bak.csv").read_bytes() == original
    assert registry.names == ["LB_tga_001", "poly9"]
    assert registry.get("poly9").metadata["batch"] == "B2"


def test_csv_keeps_byte_order_mark_and_windows_line_endings(tmp_path):
    path = tmp_path / "registry.csv"
    path.write_bytes(b"\xef\xbb\xbfname,instrument_name\r\nLB_tga_001,tga\r\n")

    append_rows(str(path), NEW_ROWS[["name", "instrument_name"]])

    assert path.read_bytes() == b"\xef\xbb\xbfname,instrument_name\r\nLB_tga_001,tga\r\npoly9,tga\r\n"


def test_xlsx_append_keeps_formatting_and_other_sheets(tmp_path):
    path = tmp_path / "registry.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "instrument_name", "sample_id", "run_id"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append(["LB_tga_001", "tga", "S001", "R001"])
    sheet.column_dimensions["A"].width = 30
    workbook.create_sheet("Notes").append(["keep me"])
    workbook.save(path)

    registry = append_rows(str(path), NEW_ROWS)

    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    assert [c.value for c in sheet[1]] == ["name", "instrument_name", "sample_id", "run_id", "raw_data_filename", "batch"]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.column_dimensions["A"].width == 30
    assert [c.value for c in sheet[3]] == ["poly9", "tga", "S009", "R001", "poly9", "B2"]
    assert workbook["Notes"]["A1"].value == "keep me"
    assert (tmp_path / "registry.bak.xlsx").exists()
    assert registry.names == ["LB_tga_001", "poly9"]


def test_locked_file_is_left_unchanged(tmp_path, monkeypatch):
    path = tmp_path / "registry.csv"
    original = "name,instrument_name\nLB_tga_001,tga\n"
    path.write_text(original, encoding="utf8")

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(PermissionError, match="Close it .* save again"):
        append_rows(str(path), NEW_ROWS)

    assert path.read_text(encoding="utf8") == original
    assert not (tmp_path / "registry.tmp.csv").exists()


def test_unsupported_format_is_refused(tmp_path):
    path = tmp_path / "registry.parquet"
    pd.DataFrame([{"name": "a", "instrument_name": "tga"}]).to_parquet(path)

    with pytest.raises(ValueError, match=".csv and .xlsx"):
        append_rows(str(path), NEW_ROWS)
