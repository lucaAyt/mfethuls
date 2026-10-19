"""The experiment registry (CSV/XLSX): one :class:`Experiment` per row, loading, validation and appending rows."""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, Iterator, List, Optional

import csv
import logging
import os
import shutil
from copy import copy

import pandas as pd

from .registry_validator import RegistryValidator
from .settings import get_settings


logger = logging.getLogger(__name__)


def resolve_registry_path(path: Optional[str] = None) -> str:
    """Resolve the registry path from an explicit argument or environment defaults."""

    if path:
        return os.path.abspath(path)

    registry_path = get_settings().registry_path
    if registry_path:
        return os.path.abspath(registry_path)

    raise ValueError(
        "No registry path provided. Pass a path explicitly or set PATH_TO_REGISTRY."
    )


def _infer_measurement_profile_from_text(text: Optional[str]) -> Optional[str]:
    """Infer a measurement profile from free-text description.

    This produces a raw registry profile value (not canonical).
    The same profile names are used for rheometer and DMA where the intent is
    comparable (frequency, strain, or temperature sweep).
    Canonicalization happens later in the parser layer.
    """

    if text is None:
        return None

    value = str(text).strip().lower()
    if not value:
        return None

    if any(token in value for token in ("freq", "frequency", "oscill")):
        return "oscillatory_frequency_sweep"
    if any(token in value for token in ("temp", "temperature")):
        return "oscillatory_temperature_sweep"
    if any(token in value for token in ("strain", "amplitude")):
        return "oscillatory_strain_sweep"
    if any(token in value for token in ("flow", "viscos", "shear")):
        return "flow_curve"

    return None


@dataclass
class Experiment:
    """Represents a single experiment definition.

    This is an abstract description, not the raw data itself. It ties together
    a human-friendly name (e.g. "CL_uv"), the instrument configuration name,
    and optional identifiers (S###, R###).

    ``experiment_id`` is a system-assigned UUID (12-char hex) set at first
    ingest via the manifest backend — it is never user-provided.

    ``raw_data_filename`` is the stem of the raw data file as the experimentalist
    named it at the instrument (e.g. "chitosan_jan15"). Defaults to ``name``
    when not declared in the registry.
    """

    name: str
    instrument_name: Optional[str]
    experiment_id: Optional[str] = None
    raw_data_filename: Optional[str] = None
    sample_id: Optional[str] = None
    run_id: Optional[str] = "R001"
    metadata: Dict[str, Any] = field(default_factory=dict)

    def __post_init__(self) -> None:
        self.sample_id = RegistryValidator.validate_sample_id(self.sample_id)
        self.run_id = RegistryValidator.validate_run_id(self.run_id)
        if self.instrument_name is None:
            return
        self.instrument_name = str(self.instrument_name).strip().casefold()


def _normalize_optional_str(value: Any) -> Optional[str]:
    """Normalize optional text fields coming from pandas records.

    Treats NaN / empty strings as ``None`` and returns a stripped string
    otherwise.
    """

    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


@dataclass
class RegistryRowResult:
    """Outcome of parsing a single registry spreadsheet row."""

    experiment: Optional[Experiment]
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def experiment_from_registry_record(rec: dict) -> RegistryRowResult:
    """Build an Experiment from a registry row without registering it."""

    errors: List[str] = []
    warnings: List[str] = []

    name = _normalize_optional_str(rec.get("name"))
    if name is None:
        errors.append("missing required field: name")
        return RegistryRowResult(experiment=None, errors=errors, warnings=warnings)

    raw_data_filename = _normalize_optional_str(rec.get("raw_data_filename")) or name

    description = _normalize_optional_str(rec.get("description"))
    explicit_measurement_profile = _normalize_optional_str(rec.get("measurement_profile"))
    sample_id = _normalize_optional_str(rec.get("sample_id"))
    raw_run_id = _normalize_optional_str(rec.get("run_id"))
    run_id = raw_run_id or "R001"
    instrument_name = _normalize_optional_str(rec.get("instrument_name"))

    measurement_profile = explicit_measurement_profile or _infer_measurement_profile_from_text(
        description
    )

    if instrument_name is None:
        warnings.append(
            "missing instrument_name; row is visible in the registry but cannot be analysed yet"
        )

    metadata: Dict[str, Any] = {}
    _known_keys = {
        "name",
        "raw_data_filename",
        "instrument_name",
        "sample_id",
        "run_id",
        "description",
        "measurement_profile",
    }
    for key, value in rec.items():
        if key in _known_keys:
            continue
        metadata[key] = value

    if description is not None:
        metadata["description"] = description
    if measurement_profile is not None:
        metadata["registry_measurement_profile"] = measurement_profile

    try:
        exp = Experiment(
            name=name,
            instrument_name=instrument_name,
            raw_data_filename=raw_data_filename,
            sample_id=sample_id,
            run_id=run_id,
            metadata=metadata,
        )
    except ValueError as exc:  # malformed sample_id / run_id
        errors.append(str(exc))
        return RegistryRowResult(experiment=None, errors=errors, warnings=warnings)
    return RegistryRowResult(experiment=exp, errors=errors, warnings=warnings)


def read_tabular_content(path: str) -> pd.DataFrame:
    """Read a CSV/XLSX file into a DataFrame."""

    _, ext = os.path.splitext(path.lower())
    if ext in {".xlsx", ".xls"}:
        df = pd.read_excel(path)
    elif ext in {".parquet", ".pq"}:
        df = pd.read_parquet(path)
    else:
        df = pd.read_csv(path)
    return df


@dataclass
class RegistryProblem:
    """A registry row that was skipped (``errors``) or loaded with ``warnings``."""

    name: Optional[str]
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def skipped(self) -> bool:
        return bool(self.errors)


class Registry:
    """The experiments described by one registry spreadsheet.

    Load it with :meth:`load` (from a CSV/XLSX/Parquet file, by default
    ``PATH_TO_REGISTRY``) or build it from a table with :meth:`from_frame`, then
    pass it to the functions that need it. Nothing is kept globally: loading the
    file again always reflects its current content.

    - ``frame``: the table as read, every row included
    - ``problems``: rows that were skipped or loaded with warnings, and why
    - ``get(name)``, ``name in registry``, iteration and ``names`` for the valid rows
    """

    def __init__(
        self,
        experiments: Iterable[Experiment] = (),
        *,
        frame: Optional[pd.DataFrame] = None,
        path: Optional[str] = None,
        problems: Optional[List[RegistryProblem]] = None,
    ) -> None:
        self._experiments: Dict[str, Experiment] = {exp.name: exp for exp in experiments}
        self.frame = frame if frame is not None else pd.DataFrame()
        self.path = path
        self.problems = problems or []

    @classmethod
    def load(
        cls,
        path: Optional[str] = None,
        *,
        check_data_paths: bool = False,
        data_root: Optional[str] = None,
    ) -> "Registry":
        """Read the registry file at ``path`` (default ``PATH_TO_REGISTRY``)."""

        resolved = resolve_registry_path(path)
        return cls.from_frame(
            read_tabular_content(resolved),
            path=resolved,
            check_data_paths=check_data_paths,
            data_root=data_root,
        )

    @classmethod
    def from_frame(
        cls,
        df: pd.DataFrame,
        *,
        path: Optional[str] = None,
        check_data_paths: bool = False,
        data_root: Optional[str] = None,
    ) -> "Registry":
        """Build a registry from a table with the registry columns.

        Required columns: ``name`` and ``instrument_name``. Optional:
        ``raw_data_filename`` (stem of the raw file; defaults to ``name``),
        ``sample_id`` (e.g. S001), ``run_id`` (e.g. R001, default R001),
        ``measurement_profile``, ``description`` (used to infer a profile when
        none is given). Any other column is kept in ``Experiment.metadata``.

        Invalid rows are skipped and recorded in ``problems``; both skipped rows
        and warnings are also logged.
        """

        source = repr(path) if path else "the given table"
        missing = {"name", "instrument_name"}.difference(df.columns)
        if missing:
            raise ValueError(f"Experiment registry at {source} is missing required columns: {sorted(missing)}")
        if "raw_data_filename" not in df.columns:
            logger.warning(
                "Registry at %s has no 'raw_data_filename' column — defaulting to 'name' for each row.", source
            )

        validator = RegistryValidator()
        experiments: List[Experiment] = []
        problems: List[RegistryProblem] = []
        for rec in df.to_dict(orient="records"):
            name = _normalize_optional_str(rec.get("name"))
            row_result = experiment_from_registry_record(rec)
            errors, warnings = list(row_result.errors), list(row_result.warnings)
            if row_result.experiment is not None and not errors:
                extra_errors, extra_warnings = validator.validate_experiment_details(
                    row_result.experiment,
                    check_data_paths=check_data_paths,
                    data_root=data_root,
                )
                errors.extend(extra_errors)
                warnings.extend(message for _, message in extra_warnings)

            if errors:
                logger.warning("Skipping registry row %r: %s", name, "; ".join(errors))
            for message in warnings:
                logger.warning("Registry row %r: %s", name, message)
            if errors or warnings:
                problems.append(RegistryProblem(name=name, errors=errors, warnings=warnings))
            if not errors and row_result.experiment is not None:
                experiments.append(row_result.experiment)

        return cls(experiments, frame=df, path=path, problems=problems)

    def get(self, name: str) -> Experiment:
        """The experiment called ``name``; ``KeyError`` if it is unknown or its row was skipped."""

        try:
            return self._experiments[name]
        except KeyError:
            where = f" at {self.path!r}" if self.path else ""
            raise KeyError(
                f"Unknown experiment name: {name!r}. It is not in the registry{where}, "
                "or its row was skipped as invalid (see Registry.problems)."
            ) from None

    def for_samples(self, sample_ids: Iterable[str]) -> List[Experiment]:
        """Experiments of the given sample ids, in registry order."""

        wanted = {str(sample_id).strip() for sample_id in sample_ids}
        return [exp for exp in self if exp.sample_id in wanted]

    @property
    def names(self) -> List[str]:
        return list(self._experiments)

    def __contains__(self, name: object) -> bool:
        return name in self._experiments

    def __iter__(self) -> Iterator[Experiment]:
        return iter(self._experiments.values())

    def __len__(self) -> int:
        return len(self._experiments)

    def __repr__(self) -> str:
        source = self.path or "table"
        return f"Registry({source!r}, {len(self)} experiments, {len(self.problems)} problems)"


def _cell(value: Any) -> Any:
    """A value as written to the registry file: missing values become empty cells."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def _append_csv(path: str, rows: pd.DataFrame, target: str) -> None:
    # Existing lines are copied as text so their values keep their exact format.
    with open(path, encoding="utf-8-sig", newline="") as handle:
        existing = list(csv.reader(handle))
    with open(path, "rb") as handle:
        raw = handle.read()
    bom = raw.startswith(b"\xef\xbb\xbf")
    line_end = "\r\n" if b"\r\n" in raw else "\n"  # keep the file's own line endings
    header = existing[0] if existing else []
    header = header + [column for column in rows.columns if column not in header]
    with open(target, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator=line_end)
        writer.writerow(header)
        for line in existing[1:]:
            writer.writerow(line + [""] * (len(header) - len(line)))
        for record in rows.to_dict(orient="records"):
            writer.writerow(["" if _cell(record.get(column)) is None else record[column] for column in header])


def _append_xlsx(path: str, rows: pd.DataFrame, target: str) -> None:
    # openpyxl edits the workbook in place: formatting, widths and other sheets are kept.
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]  # the sheet Registry.load reads
    header = [cell.value for cell in sheet[1]]
    while header and header[-1] is None:
        header.pop()
    for column in rows.columns:
        if column not in header:
            header.append(column)
            new_cell = sheet.cell(row=1, column=len(header), value=column)
            if len(header) > 1:
                new_cell.font = copy(sheet.cell(row=1, column=len(header) - 1).font)
    for record in rows.to_dict(orient="records"):
        sheet.append([_cell(record.get(column)) for column in header])
    workbook.save(target)


def append_rows(path: Optional[str], rows: pd.DataFrame) -> "Registry":
    """Append ``rows`` to the registry file at ``path`` (default ``PATH_TO_REGISTRY``).

    Existing rows and columns are kept as they are; columns the file lacks are added
    at the end, and helper columns starting with ``_`` are dropped. The previous
    version is kept next to the file as ``<name>.bak<ext>``. The new file is written
    to a temporary file first and swapped in, so a failure leaves the registry
    unchanged. Returns the reloaded registry.
    """

    path = resolve_registry_path(path)
    rows = rows[[column for column in rows.columns if not str(column).startswith("_")]]
    base, ext = os.path.splitext(path)
    writers = {".csv": _append_csv, ".xlsx": _append_xlsx}
    if ext.lower() not in writers:
        raise ValueError(f"Cannot add rows to {path!r}: only .csv and .xlsx registries can be written.")

    temporary = f"{base}.tmp{ext}"
    try:
        writers[ext.lower()](path, rows, temporary)
        shutil.copy2(path, f"{base}.bak{ext}")
        os.replace(temporary, path)
    except PermissionError as exc:
        raise PermissionError(
            f"Could not save {os.path.basename(path)}: it is open in another program. "
            "Close it (e.g. in Excel) and save again. The registry was not changed."
        ) from exc
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    logger.info("Added %d row(s) to the registry at %s", len(rows), path)
    return Registry.load(path)
